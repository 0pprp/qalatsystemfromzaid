"""Build the Qalaa API registry Excel workbook from enriched JSON-serializable rows."""

from __future__ import annotations

from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=18, color="1F4E79")
SUB_FONT = Font(name="Calibri", size=11, color="334155")
CELL_FONT = Font(name="Calibri", size=10)
AR_FONT = Font(name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
ZEBRA = PatternFill("solid", fgColor="F8FAFC")
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
FILL_YELLOW = PatternFill("solid", fgColor="FFF2CC")
FILL_RED = PatternFill("solid", fgColor="F8D7DA")
FILL_BLUE = PatternFill("solid", fgColor="D6EAF8")
FILL_GRAY = PatternFill("solid", fgColor="EEF2F6")
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_RTL = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)

ALL_COLUMNS = [
    ("endpointId", "Endpoint ID", 14),
    ("service", "Backend Service", 22),
    ("project", "Project", 22),
    ("controller", "Controller", 34),
    ("action", "Action", 28),
    ("method", "HTTP Method", 14),
    ("route", "Full Route", 42),
    ("normalizedRoute", "Normalized Route", 42),
    ("businessDomain", "Business Domain", 22),
    ("businessFunction", "Business Function", 32),
    ("arabicBusinessDescription", "Arabic Business Description", 36),
    ("technicalDescription", "Technical Description", 42),
    ("authRequired", "Auth Required", 14),
    ("authType", "Auth Type", 32),
    ("allowedRoles", "Allowed Roles", 28),
    ("internalKey", "Internal Key", 28),
    ("branchCentralScope", "Branch/Central Scope", 28),
    ("calledByApps", "Called By Apps", 36),
    ("calledByFiles", "Called By Files", 48),
    ("downstreamService", "Downstream Service", 22),
    ("downstreamEndpoint", "Downstream Endpoint", 40),
    ("repositoryService", "Repository/Service", 36),
    ("databaseTable", "Database/Table", 36),
    ("status", "Status", 16),
    ("duplicateGroupId", "Duplicate Group ID", 28),
    ("businessFunctionEndpointCount", "Business Function Endpoint Count", 16),
    ("duplicateClassification", "Duplicate Classification", 28),
    ("duplicateRiskScore", "Duplicate Risk Score", 16),
    ("potentialReplacement", "Potential Replacement", 36),
    ("notes", "Notes", 40),
    ("sourceFile", "Source File", 48),
    ("sourceLine", "Source Line", 12),
]

APP_COLUMNS = [
    ("feature", "Feature", 28),
    ("screen", "Screen/Module", 28),
    ("method", "HTTP Method", 14),
    ("normalizedRoute", "Endpoint", 42),
    ("service", "Backend Service", 22),
    ("technicalDescription", "Purpose", 40),
    ("authType", "Auth", 28),
    ("calledFrom", "Called From File", 48),
    ("businessFunction", "Business Function", 32),
    ("duplicateGroupId", "Duplicate Group", 24),
    ("businessFunctionEndpointCount", "Duplicate Count", 14),
    ("notes", "Notes", 36),
]

BE_COLUMNS = ALL_COLUMNS

DUP_COLUMNS = [
    ("duplicateGroupId", "Duplicate Group ID", 28),
    ("businessFunction", "Business Function", 32),
    ("endpointCount", "Endpoint Count", 14),
    ("endpoints", "Endpoints", 70),
    ("services", "Services", 36),
    ("consumers", "Consumers", 36),
    ("classification", "Classification", 28),
    ("riskScore", "Risk Score", 12),
    ("whyMultipleExist", "Why Multiple Exist", 50),
    ("recommendedCanonicalEndpoint", "Recommended Canonical Endpoint", 42),
    ("canConsolidate", "Can Consolidate?", 16),
    ("recommendedAction", "Recommended Action", 46),
    ("migrationRisk", "Migration Risk", 18),
]

CHAIN_COLUMNS = [
    ("feature", "Feature", 36),
    ("client", "Client", 28),
    ("clientCall", "Client Call", 46),
    ("branch", "Branch", 40),
    ("internalCall", "Internal Call", 42),
    ("gateway", "Gateway", 36),
    ("service", "Service", 32),
    ("table", "Table", 36),
    ("notes", "Notes", 50),
]

ORPHAN_COLUMNS = [
    ("method", "HTTP Method", 12),
    ("normalizedRoute", "Endpoint", 42),
    ("service", "Service", 22),
    ("businessFunction", "Business Function", 32),
    ("technicalDescription", "Last-known purpose", 40),
    ("staticConsumers", "Static Consumers Found", 18),
    ("duplicateGroupId", "Duplicate Group", 24),
    ("duplicateRiskScore", "Risk", 10),
    ("recommendation", "Recommendation", 40),
    ("evidence", "Evidence", 50),
]

CANON_COLUMNS = [
    ("businessFunction", "Business Function", 32),
    ("current", "Current Endpoints", 70),
    ("publicCanonical", "Preferred Canonical Public Endpoint", 46),
    ("internalCanonical", "Preferred Internal Endpoint", 42),
    ("owner", "Owner Service", 22),
    ("clients", "Client Apps", 36),
    ("legacy", "Legacy Endpoint(s)", 42),
    ("futureRule", "Future Rule", 50),
]


def _style_header(ws: Worksheet, ncols: int):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"
    ws.row_dimensions[1].height = 22
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN


def _write_table(ws: Worksheet, columns: list[tuple], rows: list[dict], arabic_cols: set[str] | None = None):
    arabic_cols = arabic_cols or set()
    for i, (key, title, width) in enumerate(columns, 1):
        ws.cell(1, i, title)
        ws.column_dimensions[get_column_letter(i)].width = width
    _style_header(ws, len(columns))
    for r_idx, row in enumerate(rows, 2):
        for c_idx, (key, _title, _w) in enumerate(columns, 1):
            val = row.get(key, "")
            if isinstance(val, list):
                val = "; ".join(str(x) for x in val)
            if isinstance(val, bool):
                val = "Yes" if val else "No"
            cell = ws.cell(r_idx, c_idx, val)
            cell.font = AR_FONT if key in arabic_cols else CELL_FONT
            cell.alignment = WRAP_RTL if key in arabic_cols else WRAP
            cell.border = THIN
            if r_idx % 2 == 0:
                cell.fill = ZEBRA
            if key == "duplicateRiskScore" and isinstance(row.get(key), int):
                score = row[key]
                if score >= 4:
                    cell.fill = FILL_RED
                elif score == 3:
                    cell.fill = FILL_YELLOW
                else:
                    cell.fill = FILL_GREEN
            if key == "duplicateClassification":
                text = str(row.get(key) or "")
                if text.startswith("E."):
                    cell.fill = FILL_RED
                elif text.startswith("D.") or text.startswith("F."):
                    cell.fill = FILL_YELLOW
                elif text.startswith("B."):
                    cell.fill = FILL_BLUE
        ws.row_dimensions[r_idx].height = 32
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    ws.sheet_view.showGridLines = False


def _join(values) -> str:
    if isinstance(values, list):
        return "; ".join(str(v) for v in values if v)
    return str(values or "")


def build_workbook(data: dict, out_path: Path) -> None:
    wb = Workbook()
    iraq = timezone(timedelta(hours=3))
    generated = data["meta"]["generatedAt"]
    commit = data["meta"]["commit"]

    # 00_Summary
    ws = wb.active
    ws.title = "00_Summary"
    ws.sheet_view.rightToLeft = False
    ws["A1"] = "Qalaa API / Endpoint Architecture Registry"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws["A2"] = (
        f"Generated {generated} (Iraq time)  |  git {commit}  |  "
        "ANALYSIS ONLY — no runtime behavior changed"
    )
    ws["A2"].font = SUB_FONT
    ws.merge_cells("A2:F2")

    summary = data["summary"]
    metrics = [
        ("Total endpoints", summary["totalEndpoints"]),
        ("GET", summary["byMethod"].get("GET", 0)),
        ("POST", summary["byMethod"].get("POST", 0)),
        ("PUT", summary["byMethod"].get("PUT", 0)),
        ("PATCH", summary["byMethod"].get("PATCH", 0)),
        ("DELETE", summary["byMethod"].get("DELETE", 0)),
        ("SIGNALR", summary["byMethod"].get("SIGNALR", 0)),
        ("Business functions", summary["businessFunctionCount"]),
        ("Unique (A)", summary["classificationCounts"].get("A. UNIQUE", 0)),
        ("Legitimate bridge (B)", summary["classificationCounts"].get("B. LEGITIMATE BRIDGE", 0)),
        ("Role-specific (C)", summary["classificationCounts"].get("C. ROLE-SPECIFIC", 0)),
        ("Legacy (D)", summary["classificationCounts"].get("D. LEGACY", 0)),
        ("True duplicate groups (E)", summary["trueDuplicateGroupCount"]),
        ("Review required (F)", summary["classificationCounts"].get("F. UNKNOWN / REVIEW REQUIRED", 0)),
        ("High-risk duplicate groups (score ≥ 4)", summary["highRiskGroupCount"]),
        ("Endpoints with no detected client consumer", summary["noClientConsumerCount"]),
        ("Endpoints without auth", summary["unauthenticatedCount"]),
        ("Internal endpoints", summary["internalCount"]),
        ("Mobile-only consumed", summary["mobileOnlyCount"]),
        ("Web-only consumed", summary["webOnlyCount"]),
        ("Shared web+mobile", summary["sharedCount"]),
    ]
    ws["A4"] = "Metric"
    ws["B4"] = "Value"
    _style_header(ws, 2)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 18
    for i, (k, v) in enumerate(metrics, 5):
        ws.cell(i, 1, k).font = CELL_FONT
        ws.cell(i, 2, v).font = Font(name="Calibri", bold=True, size=11)
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2).border = THIN
        if i % 2 == 0:
            ws.cell(i, 1).fill = ZEBRA
            ws.cell(i, 2).fill = ZEBRA

    ws["D4"] = "Backend service"
    ws["E4"] = "Endpoints"
    for col in (4, 5):
        ws.cell(4, col).fill = HEADER_FILL
        ws.cell(4, col).font = HEADER_FONT
        ws.cell(4, col).border = THIN
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 14
    for i, (svc, n) in enumerate(summary["byService"].items(), 5):
        ws.cell(i, 4, svc).border = THIN
        ws.cell(i, 5, n).border = THIN

    ws["D10"] = "Client app"
    ws["E10"] = "Consumed endpoints"
    ws["D10"].fill = HEADER_FILL
    ws["E10"].fill = HEADER_FILL
    ws["D10"].font = HEADER_FONT
    ws["E10"].font = HEADER_FONT
    row = 11
    for app, n in summary["byClientApp"].items():
        ws.cell(row, 4, app).border = THIN
        ws.cell(row, 5, n).border = THIN
        row += 1

    note_row = max(28, row + 2)
    ws.cell(note_row, 1, "How to read duplication").font = Font(name="Calibri", bold=True, size=13, color="1F4E79")
    notes = [
        "A UNIQUE — one implementation. Keep.",
        "B LEGITIMATE BRIDGE — client/branch entry + trusted internal or gateway proxy. Required by architecture.",
        "C ROLE-SPECIFIC — same capability, different auth/role (manager vs employee vs filter vs delegate).",
        "D LEGACY — older alias still present; prefer modern route for new work.",
        "E TRUE DUPLICATE — same responsibility AND same trust boundary. Do not add more.",
        "F REVIEW REQUIRED — not enough static evidence.",
        "Governance: one canonical public entry point per responsibility AND trust boundary.",
        "Do not delete an endpoint only because a similar one exists. Check trust boundaries first.",
        "NO STATIC CONSUMER FOUND ≠ unused. It may be called by production builds, Postman, or other provinces.",
        "Secrets (connection strings, JWT keys, gateway keys) are excluded from this workbook.",
    ]
    for i, line in enumerate(notes, note_row + 1):
        ws.cell(i, 1, line).font = SUB_FONT
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)

    # 01_All_Endpoints
    ws_all = wb.create_sheet("01_All_Endpoints")
    all_rows = []
    for r in data["endpoints"]:
        item = dict(r)
        item["calledByApps"] = _join(r.get("calledByApps"))
        item["calledByFiles"] = _join(r.get("calledByFiles"))
        item["authRequired"] = "Yes" if r.get("authRequired") else "No"
        all_rows.append(item)
    _write_table(ws_all, ALL_COLUMNS, all_rows, {"arabicBusinessDescription", "allowedRoles"})

    # App sheets
    app_sheet_names = data["meta"]["appSheets"]
    for app, sheet_name in app_sheet_names.items():
        rows = []
        for r in data["endpoints"]:
            files = [f for f in r.get("calledByFiles", []) if app.replace("_application", "") in f or f.startswith(app) or f"/{app}/" in f.replace("\\", "/") or f.startswith(app)]
            if app not in r.get("calledByApps", []):
                continue
            called = [f for f in r.get("calledByFiles", []) if app in f or True]
            # Keep files that belong to this app
            app_files = [f for f in r.get("calledByFiles", []) if f.replace("\\", "/").startswith(app) or f"/{app}/" in f.replace("\\", "/")]
            rows.append(
                {
                    "feature": r.get("businessDomain"),
                    "screen": app_files[0].split("/")[-1] if app_files else "",
                    "method": r["method"],
                    "normalizedRoute": r["normalizedRoute"],
                    "service": r["service"],
                    "technicalDescription": r.get("technicalDescription"),
                    "authType": r.get("authType"),
                    "calledFrom": _join(app_files or r.get("calledByFiles")),
                    "businessFunction": r.get("businessFunction"),
                    "duplicateGroupId": r.get("duplicateGroupId"),
                    "businessFunctionEndpointCount": r.get("businessFunctionEndpointCount"),
                    "notes": r.get("notes"),
                }
            )
        ws_app = wb.create_sheet(sheet_name[:31])
        _write_table(ws_app, APP_COLUMNS, rows)

    # Backend sheets
    for svc, sheet_name in data["meta"]["serviceSheets"].items():
        rows = [item for item in all_rows if item["service"] == svc]
        ws_be = wb.create_sheet(sheet_name[:31])
        _write_table(ws_be, BE_COLUMNS, rows, {"arabicBusinessDescription", "allowedRoles"})

    # 90 duplicates
    dup_rows = []
    for g in data["duplicateGroups"]:
        item = dict(g)
        item["endpoints"] = " | ".join(g.get("endpoints") or [])
        item["services"] = _join(g.get("services"))
        item["consumers"] = _join(g.get("consumers"))
        item["canConsolidate"] = "Yes" if g.get("canConsolidate") else "No"
        dup_rows.append(item)
    ws_dup = wb.create_sheet("90_Duplicate_Analysis")
    _write_table(ws_dup, DUP_COLUMNS, dup_rows)

    ws_ch = wb.create_sheet("91_Call_Chains")
    _write_table(ws_ch, CHAIN_COLUMNS, data["callChains"])

    orphan_rows = []
    for r in data["orphans"]:
        orphan_rows.append(
            {
                "method": r["method"],
                "normalizedRoute": r["normalizedRoute"],
                "service": r["service"],
                "businessFunction": r["businessFunction"],
                "technicalDescription": r.get("technicalDescription"),
                "staticConsumers": "NO STATIC CONSUMER FOUND",
                "duplicateGroupId": r.get("duplicateGroupId"),
                "duplicateRiskScore": r.get("duplicateRiskScore"),
                "recommendation": r.get("recommendation"),
                "evidence": f"{r.get('sourceFile')}:{r.get('sourceLine')}",
            }
        )
    ws_or = wb.create_sheet("92_Possible_Orphans")
    _write_table(ws_or, ORPHAN_COLUMNS, orphan_rows)

    ws_st = wb.create_sheet("93_API_Standards")
    _write_standards(ws_st, generated, commit)

    ws_ca = wb.create_sheet("94_Canonical_Map")
    _write_table(ws_ca, CANON_COLUMNS, data["canonicalMap"])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def _write_standards(ws: Worksheet, generated: str, commit: str) -> None:
    ws["A1"] = "Qalaa API naming and governance standard"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")
    ws["A2"] = f"Generated {generated}  |  {commit}"
    ws["A2"].font = SUB_FONT
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 100
    rows = [
        ("Collections", "GET /api/sales-requests    POST /api/sales-requests"),
        ("Resource", "GET /api/sales-requests/{id}    PUT /api/sales-requests/{id}    DELETE /api/sales-requests/{id}"),
        ("State / action", "POST /api/sales-requests/{id}/assign    POST /api/sales-exceptions/{id}/approve    POST /api/sales-exceptions/{id}/reject"),
        ("Internal S2S", "POST /api/internal/sales-exceptions    POST /api/internal/complaints"),
        ("Case", "lowercase kebab-case path segments. Plural resource nouns."),
        ("Verbs", "Do not put verbs in the collection name (avoid CreateSalesRequest, GetAllSalesRequestsNew). Use HTTP method + resource, or a sub-resource action when the operation is not CRUD."),
        ("Prefix", "Public HTTP APIs live under /api/. SignalR hubs live under /hubs/."),
        ("Internal prefix", "/api/internal/ is only for trusted service-to-service calls authenticated by X-Sales-Gateway-Key (or equivalent). Browsers and mobile apps must never call it."),
        ("Avoid", "createSalesRequest; GetAllSalesRequestsNew; SalesManager/CreateExceptionRequest2; mixed singular/plural; duplicate prefixes (/api/api); encoding query params in the path with &&."),
        ("Versioning", "When a breaking change is required, add /api/v2/... rather than a parallel unversioned alias. Keep v1 until clients migrate."),
        ("Pagination", "Prefer ?page=&pageSize= with a stable envelope { items, total, page, pageSize }. Do not invent a new paging shape per controller."),
        ("Errors", "Return JSON { message } (Arabic user-facing text is fine). Use 401 for auth, 403 for role, 409 for business conflict, 502 when a bridge cannot reach the next hop."),
        ("Auth", "JWT Bearer for people. X-Sales-Gateway-Key only on /api/internal/* and explicit gateway-key branch endpoints. Delegate apps: X-Async-Id (+ optional X-Delegate-Id). Portal: encrypted query token."),
        ("Idempotency", "POSTs that create money/sales/exceptions should accept Idempotency-Key or reject duplicates by natural key (already done for active exceptions)."),
        ("Trust boundary rule", "One business capability should have one canonical public entry point per trust boundary. A branch public API plus a central internal intake is TWO endpoints on purpose."),
        ("Do not globally collapse", "Do not force Flutter, FE, and internal bridges onto a single physical URL. Collapse only when A, B, and C share the same caller type AND the same trust boundary."),
        ("BEFORE ADDING ANY NEW ENDPOINT", "1. Search Qalaa_API_Endpoint_Registry.xlsx  2. Search code by Business Function  3. Search DuplicateGroupId  4. Check whether an existing endpoint can be extended  5. Check service ownership  6. Check whether a bridge already exists  7. Add a new route only if existing routes cannot support the requirement cleanly  8. Document justification  9. Update this registry"),
        ("Demo vs production", "Demo hosts (e.g. :8080 BE_Company, :8081 delegate, :5280 gateway local) must not become extra public contracts. Same route shape everywhere; only the base URL changes."),
        ("City addressing", "Prefer ?cityValue= on central list endpoints. Path {cityValue} is acceptable for single-resource fan-out but do not keep both evaluate variants."),
        ("Legacy ERP names", "Existing Customers_GetAll / Delegates_Create routes stay until FE_Company is migrated. New work must not copy that style."),
    ]
    ws.cell(4, 1, "Topic").fill = HEADER_FILL
    ws.cell(4, 2, "Rule").fill = HEADER_FILL
    ws.cell(4, 1).font = HEADER_FONT
    ws.cell(4, 2).font = HEADER_FONT
    for i, (k, v) in enumerate(rows, 5):
        ws.cell(i, 1, k).font = Font(name="Calibri", bold=True, size=10)
        ws.cell(i, 2, v).alignment = WRAP
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2).border = THIN
        ws.row_dimensions[i].height = 48
    ws.freeze_panes = "A5"
